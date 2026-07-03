"""Lectura del diccionario de municipios del INE (diccionario{YY}.xlsx).

Hoja única con fila de título, cabecera `CODAUTO, CPRO, CMUN, DC, NOMBRE` y una
fila por municipio. El código INE de 5 dígitos es CPRO (2) + CMUN (3).
"""
import openpyxl

from .modelo import Municipio
from .territorios import CCAA_POR_CODAUTO, PROVINCIA_POR_CPRO


def leer_municipios(xlsx_path) -> list[Municipio]:
    """Devuelve los municipios normalizados desde el xlsx del INE."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    municipios: list[Municipio] = []
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        codauto = str(row[0]).strip()
        if not codauto.isdigit():
            continue  # fila de título o cabecera
        codauto = codauto.zfill(2)
        cpro = str(row[1]).strip().zfill(2)
        cmun = str(row[2]).strip().zfill(3)
        iso, ccaa_nombre = CCAA_POR_CODAUTO[codauto]
        municipios.append(Municipio(
            ine=cpro + cmun,
            name=str(row[4]).strip(),
            province=cpro,
            province_name=PROVINCIA_POR_CPRO[cpro],
            ccaa_ine=codauto,
            ccaa_iso=iso,
            ccaa_name=ccaa_nombre,
            dc=int(str(row[3]).strip()),
        ))
    return municipios
