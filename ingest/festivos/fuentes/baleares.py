"""Fuente de fiestas locales de Illes Balears (Govern de les Illes Balears).

El Govern publica el calendari laboral en su portal de datos abiertos (CAIB,
ficha datos.gob.es `a04003003`) como un recurso DISTINTO POR AÑO. No hay una
sola URL: cada año tiene su propio dataset/recurso, y el formato cambia entre
años (CSV en 2024 y 2026, solo XLSX en 2025). Todos traen TODOS los ámbitos
(autonómico y local) con columnas `Illa,Àmbit,Municipi,Localitat,Data,Nom
festa`; aquí nos quedamos con las filas `Àmbit == "Local"`.

La fecha viene en formatos distintos según el recurso:

* texto catalán, p. ej. `24 de juny`, `2 d'abril` (CSV 2026);
* `DD/MM/YYYY`, p. ej. `16/08/2024` (CSV 2024);
* fecha/datetime real de Excel (XLSX 2025).

IMPORTANTE: el recurso de un año NUNCA se reutiliza para otro. Reutilizarlo
produciría datos FALSOS, porque el documento contiene las fechas reales de SU
año (que en el formato de texto catalán ni siquiera llevan el año). Para un año
sin fuente conocida se lanza `ValueError`.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import urllib.request
from pathlib import Path
from typing import Iterable

import openpyxl

from .base import FuenteLocal, FiestaLocalCruda, registrar

# Nombre del mes en catalán (normalizado, sin tildes) -> número de mes.
_MESES: dict[str, int] = {
    "gener": 1,
    "febrer": 2,
    "marc": 3,
    "abril": 4,
    "maig": 5,
    "juny": 6,
    "juliol": 7,
    "agost": 8,
    "setembre": 9,
    "octubre": 10,
    "novembre": 11,
    "desembre": 12,
}

# "24 de juny", "2 d'abril", "2 d’abril" -> (día, mes en texto).
_RE_FECHA_TEXTO = re.compile(r"^\s*(\d{1,2})\s+d[e'’]?\s*(.+?)\s*$")
# "16/08/2024" o "16-08-2024" -> (día, mes, año) numéricos.
_RE_FECHA_NUM = re.compile(r"^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{4})\s*$")

# Recurso de datos abiertos (CAIB) por año. Cada año tiene su propio dataset y
# recurso; el formato indica cómo parsear/cachear. Resuelto vía la ficha
# datos.gob.es `a04003003` (un dataset por año). Para añadir un año nuevo basta
# con dar de alta su recurso aquí.
_FUENTES_POR_ANYO: dict[int, dict[str, str]] = {
    2024: {
        "formato": "csv",
        "url": (
            "https://intranet.caib.es/opendatacataleg/dataset/"
            "9b3aae18-62c5-4030-aaf3-768aefc429c3/resource/"
            "ded8d213-e500-43c6-96c8-a4259557bde5/download/"
            "calendari-laboral-general-i-local-illes-balears-2024_2.csv"
        ),
    },
    2025: {
        # El recurso de 2025 solo se publica en XLSX (no hay CSV).
        "formato": "xlsx",
        "url": (
            "https://intranet.caib.es/opendatacataleg/dataset/"
            "aeac49dc-a6d8-47be-911d-75225a60e34f/resource/"
            "cf3686ae-5576-43d2-9cec-c2b12c792db8/download/"
            "calendari-laboral-2025.xlsx"
        ),
    },
    2026: {
        "formato": "csv",
        "url": (
            "https://intranet.caib.es/opendatacataleg/dataset/"
            "e89fb44b-67f3-4e29-affc-2df135b719e5/resource/"
            "edf08154-bbc0-4259-a254-3b0185411354/download/"
            "calendari-laboral-2026.csv"
        ),
    },
}


def _normalizar(texto: str) -> str:
    """Pasa a minúsculas y elimina tildes para casar nombres de mes."""
    sin_tildes = unicodedata.normalize("NFKD", texto)
    sin_tildes = "".join(c for c in sin_tildes if not unicodedata.combining(c))
    return sin_tildes.strip().lower()


@registrar
class Baleares(FuenteLocal):
    """Fiestas locales de Illes Balears desde el calendario laboral (CAIB)."""
    ccaa_iso = "ES-IB"
    nombre = "Illes Balears"

    def descargar(self, anyo: int) -> bytes:
        """Descarga (o lee de caché) el recurso del calendario laboral del año.

        Resuelve la URL del recurso CORRECTO de cada año en `_FUENTES_POR_ANYO`.
        La caché es POR AÑO (la clave incluye el año), de modo que el documento
        de un año nunca se confunde con el de otro. Si no hay fuente conocida
        para el año pedido, falla en vez de reutilizar otro documento.
        """
        fuente = _FUENTES_POR_ANYO.get(anyo)
        if fuente is None:
            raise ValueError(f"sin fuente de {self.ccaa_iso} para {anyo}")

        raiz = Path(__file__).resolve().parents[3]
        cache = raiz / ".cache" / f"local_ES-IB_{anyo}.{fuente['formato']}"
        if cache.exists():
            return cache.read_bytes()

        with urllib.request.urlopen(fuente["url"]) as resp:
            bruto = resp.read()

        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(bruto)
        return bruto

    def parsear(self, bruto: bytes, anyo: int) -> Iterable[FiestaLocalCruda]:
        """Convierte el recurso del año en fiestas locales (CPRO 07).

        Según el año el recurso es CSV (texto) o XLSX (binario Excel); el
        formato se deduce de `_FUENTES_POR_ANYO`. En ambos casos las columnas
        lógicas son `Illa,Àmbit,Municipi,Localitat,Data,Nom festa` y nos
        quedamos con las filas `Àmbit == "Local"`.
        """
        fuente = _FUENTES_POR_ANYO.get(anyo)
        if fuente is None:
            raise ValueError(f"sin fuente de {self.ccaa_iso} para {anyo}")

        if fuente["formato"] == "xlsx":
            filas = self._filas_xlsx(bruto)
        else:
            filas = self._filas_csv(bruto)

        for ambito, municipio, data in filas:
            if ambito != "local":
                continue  # solo fiestas locales/municipales

            municipio_nombre = municipio.strip()
            fecha = self._fecha_iso(data, anyo)
            if not fecha or not municipio_nombre:
                continue  # parseo defensivo: fila incompleta o fecha ilegible

            denominacion = (data[1] if isinstance(data, tuple) else "")
            yield FiestaLocalCruda(
                fecha=fecha,
                municipio_nombre=municipio_nombre,
                ine=None,          # se resuelve después con el Matcher (nombre -> INE)
                provincia="07",    # Illes Balears (CPRO), única provincia de la CCAA
                denominacion=denominacion or None,
            )

    @staticmethod
    def _filas_csv(bruto: bytes) -> Iterable[tuple[str, str, tuple[str, str]]]:
        """Itera el CSV devolviendo `(ambito, municipio, (fecha, nom_festa))`.

        El CSV puede venir en latin-1 (2026) o UTF-8 con BOM (2024); se prueba
        UTF-8(-sig) primero y se cae a latin-1, que nunca falla al decodificar.
        """
        try:
            texto = bruto.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = bruto.decode("latin-1")
        lector = csv.DictReader(io.StringIO(texto))
        for fila in lector:
            yield (
                (fila.get("Àmbit") or "").strip().lower(),
                (fila.get("Municipi") or "").strip(),
                ((fila.get("Data") or "").strip(), (fila.get("Nom festa") or "").strip()),
            )

    @staticmethod
    def _filas_xlsx(bruto: bytes) -> Iterable[tuple[str, str, tuple[object, str]]]:
        """Itera el XLSX devolviendo `(ambito, municipio, (fecha, nom_festa))`.

        La hoja única tiene cabecera `Illa,Àmbit,Municipi,Localitat,Data,Nom
        festa`; la columna `Data` es una fecha/datetime real de Excel.
        """
        wb = openpyxl.load_workbook(io.BytesIO(bruto), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for fila in ws.iter_rows(values_only=True):
            if not fila or len(fila) < 6 or fila[1] is None:
                continue
            ambito = str(fila[1]).strip().lower()
            if ambito == "àmbit":
                continue  # fila de cabecera
            municipio = str(fila[2]).strip() if fila[2] is not None else ""
            nom_festa = str(fila[5]).strip() if fila[5] is not None else ""
            yield (ambito, municipio, (fila[4], nom_festa))

    @classmethod
    def _fecha_iso(cls, data, anyo: int) -> str | None:
        """Convierte la fecha del recurso a ISO `YYYY-MM-DD`.

        `data` es `(valor_fecha, nom_festa)`. El valor de fecha puede ser texto
        catalán (`24 de juny`), `DD/MM/YYYY` o un `date`/`datetime` de Excel.
        """
        valor = data[0] if isinstance(data, tuple) else data

        # XLSX: la fecha ya viene como objeto date/datetime.
        if hasattr(valor, "year") and hasattr(valor, "month") and hasattr(valor, "day"):
            if valor.year != anyo:
                return None  # defensivo: el recurso debe ser del año pedido
            return f"{valor.year:04d}-{valor.month:02d}-{valor.day:02d}"

        texto = (valor or "").strip()
        if not texto:
            return None

        # Formato numérico DD/MM/YYYY (CSV 2024).
        num = _RE_FECHA_NUM.match(texto)
        if num:
            dia, mes, anyo_doc = int(num.group(1)), int(num.group(2)), int(num.group(3))
            if anyo_doc != anyo or not 1 <= mes <= 12 or not 1 <= dia <= 31:
                return None
            return f"{anyo:04d}-{mes:02d}-{dia:02d}"

        # Formato texto catalán "24 de juny" (CSV 2026): no lleva el año, se
        # toma el del recurso (que ya es el correcto por la caché por año).
        coincidencia = _RE_FECHA_TEXTO.match(texto)
        if not coincidencia:
            return None
        dia = int(coincidencia.group(1))
        mes_token = _normalizar(coincidencia.group(2))
        mes = _MESES.get(mes_token)
        if mes is None:
            # Tolerar erratas castellanas como "agosto" -> "agost".
            mes = _MESES.get(mes_token.rstrip("o"))
        if mes is None or not 1 <= dia <= 31:
            return None
        return f"{anyo:04d}-{mes:02d}-{dia:02d}"
